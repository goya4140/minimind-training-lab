#!/usr/bin/env python3
"""Build an internal methods draft, never a substitute for the final report.

Run with the bundled workspace Python. Output belongs in .report-build/.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FONT = "Arial Unicode MS"


def load_content(path: Path) -> dict:
    content = json.loads(path.read_text(encoding="utf-8"))
    if content.get("schema") != 1 or content.get("status") != "methods-draft":
        raise ValueError("This builder accepts methods-draft only, not final results")
    ids = [section["id"] for section in content["sections"]]
    if len(ids) != len(set(ids)) or not content["missing_final_sections"]:
        raise ValueError("Draft needs unique sections and explicit remaining final requirements")
    for section in content["sections"]:
        for source in section["sources"]:
            if not (ROOT / source).is_file():
                raise FileNotFoundError(source)
        if image := section.get("image"):
            actual = hashlib.sha256((ROOT / image["path"]).read_bytes()).hexdigest()
            if actual != image["sha256"]:
                raise ValueError("Evidence image checksum mismatch")
    return content


def font_style(style, size: float, *, bold: bool = False) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), FONT)


def set_cell_shading(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc: Document, data: dict) -> None:
    columns = len(data["headers"])
    values = [data["headers"], *data["rows"]]
    table = doc.add_table(rows=len(values), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = data["widths"]
    for column, proportion in zip(table.columns, widths):
        column.width = Inches(7 * proportion)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        for attribute, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            line.set(qn(f"w:{attribute}"), value)
        borders.append(line)
    table._tbl.tblPr.append(borders)
    for r, values_row in enumerate(values):
        for c, value in enumerate(values_row):
            cell = table.cell(r, c)
            cell.width = Inches(7 * widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.1
            if c > 0 and len(str(value)) < 12:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            run.font.size = Pt(10)
            run.font.bold = r == 0
            run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else "000000")
            set_cell_shading(cell, "20394D" if r == 0 else ("F1F5F8" if r % 2 == 0 else "FFFFFF"))
        no_split = OxmlElement("w:cantSplit")
        table.rows[r]._tr.get_or_add_trPr().append(no_split)
    repeat_header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(content: dict, output: Path) -> None:
    if output.exists():
        raise FileExistsError("Use a new revision filename; existing draft is preserved")
    if ".report-build" not in output.resolve().parts:
        raise ValueError("Methods drafts must stay in private .report-build, not final output")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    font_style(doc.styles["Normal"], 11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.2
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    font_style(doc.styles["Title"], 27, bold=True)
    font_style(doc.styles["Subtitle"], 16)
    font_style(doc.styles["Heading 1"], 19, bold=True)
    font_style(doc.styles["Heading 2"], 14, bold=True)
    font_style(doc.styles["Caption"], 9)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.keep_with_next = True
    # Remove inherited template border and theme overrides, including Title's rule.
    for element in list(doc.styles.element.iter(qn("w:pBdr"))):
        element.getparent().remove(element)
    for fonts in doc.styles.element.iter(qn("w:rFonts")):
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
    doc.core_properties.title = content["title"] + " " + content["subtitle"]
    doc.core_properties.subject = "MiniMind fixed implementation and methods"
    doc.core_properties.author = "MiniMind Training Lab"
    doc.add_paragraph(content["title"], "Title")
    doc.add_paragraph(content["subtitle"], "Subtitle")
    doc.add_paragraph(content["scope"])
    doc.add_paragraph(f"正式运行 {content['run_id']}")
    doc.add_paragraph(f"训练源码 {content['training_commit']}", "Caption")
    for index, topic in enumerate(content["sections"], 1):
        doc.add_page_break()
        doc.add_heading(f"{index} {topic['title']}", level=1)
        for paragraph in topic["paragraphs"]:
            doc.add_paragraph(paragraph)
        if "table" in topic:
            add_table(doc, topic["table"])
        if image := topic.get("image"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(ROOT / image["path"]), width=Inches(6.7))
            doc.add_paragraph(image["caption"], "Caption")
        source_text = "；".join(topic["sources"])
        doc.add_paragraph(f"实现与证据来源 {source_text}", "Caption")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--content", type=Path, default=ROOT / "reporting/methods-content.json")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    content = load_content(args.content)
    build(content, args.output)
    print(json.dumps({"output": str(args.output), "status": content["status"], "sections": len(content["sections"])}))


if __name__ == "__main__":
    main()
